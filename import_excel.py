#!/usr/bin/env python3
"""将 output 目录下的 Excel 文件导入到数据库中"""

import sys
import os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import load_workbook
from datetime import datetime
from server.database import init_db, SessionLocal
from server.models import Lead

# Excel 列映射（根据 exporter.py 的 COLUMNS 定义）
# 序号(1) 公司名称(2) 联系电话(3) 所在城市(4) 区县(5) 详细地址(6) 行业分类(7) 数据来源(8) 搜索词(9)
COL_MAP = {
    2: "name",
    3: "phone",
    4: "city",
    5: "district",
    6: "address",
    7: "industry",
    8: "source",
    9: "search_keyword",
}


def import_file(filepath: str):
    print(f"读取文件: {filepath}")
    wb = load_workbook(filepath, read_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(min_row=3, values_only=True))  # 跳过标题行(1)和表头(2)
    print(f"读取到 {len(rows)} 行数据")

    init_db()
    db = SessionLocal()

    added = 0
    updated = 0
    skipped = 0

    for row in rows:
        if not row or len(row) < 9:
            skipped += 1
            continue

        name = str(row[1] or "").strip()
        if not name:
            skipped += 1
            continue

        phone = str(row[2] or "").strip()
        city = str(row[3] or "").strip()
        district = str(row[4] or "").strip()
        address = str(row[5] or "").strip()
        industry = str(row[6] or "").strip()
        source = str(row[7] or "").strip()
        search_keyword = str(row[8] or "").strip()

        existing = db.query(Lead).filter(Lead.name == name).first()
        if existing:
            # 已存在：如果原来没电话但新数据有，则更新
            if not existing.phone and phone:
                existing.phone = phone
                existing.updated_at = datetime.now()
                updated += 1
            else:
                skipped += 1
        else:
            lead = Lead(
                name=name,
                phone=phone,
                city=city,
                district=district,
                address=address,
                industry=industry,
                source=source,
                search_keyword=search_keyword,
            )
            db.add(lead)
            added += 1

        # 每 500 条提交一次
        if (added + updated) % 500 == 0 and (added + updated) > 0:
            db.commit()
            print(f"  进度: 新增 {added}, 更新 {updated}, 跳过 {skipped}")

    db.commit()
    db.close()
    wb.close()

    print()
    print(f"导入完成!")
    print(f"  新增: {added} 条")
    print(f"  更新: {updated} 条")
    print(f"  跳过: {skipped} 条（重复或空行）")


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else "output/服装厂名单.xlsx"

    if not os.path.exists(target):
        print(f"文件不存在: {target}")
        sys.exit(1)

    import_file(target)
