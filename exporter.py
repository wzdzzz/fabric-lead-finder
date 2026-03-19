"""
数据导出模块 - 将采集结果导出为格式化 Excel
"""

import os
import logging
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import OUTPUT_FILENAME
from paths import OUTPUT_DIR

logger = logging.getLogger(__name__)

# Excel 表头: (显示名, 列宽, 数据字段)
COLUMNS = [
    ("序号", 6, None),
    ("公司名称", 35, "name"),
    ("联系电话", 22, "phone"),
    ("所在城市", 10, "city"),
    ("区县", 10, "district"),
    ("详细地址", 45, "address"),
    ("行业分类", 20, "industry"),
    ("数据来源", 10, "source"),
    ("搜索词", 12, "search_keyword"),
]


def deduplicate(companies):
    """按公司名称去重，优先保留有电话的记录"""
    seen = {}
    for c in companies:
        name = c.get("name", "").strip()
        if not name:
            continue
        # 如果已存在，优先保留有电话的
        if name in seen:
            if not seen[name].get("phone") and c.get("phone"):
                seen[name] = c
        else:
            seen[name] = c

    unique = list(seen.values())
    removed = len(companies) - len(unique)
    if removed > 0:
        logger.info(f"去重: 移除 {removed} 条重复记录")
    return unique


def export_to_excel(companies, filename=None):
    """将企业列表导出为格式化的 Excel 文件"""
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    if filename is None:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        name, ext = os.path.splitext(OUTPUT_FILENAME)
        filename = f"{name}_{timestamp}{ext}"

    filepath = os.path.join(OUTPUT_DIR, filename)

    # 去重
    companies = deduplicate(companies)

    # 统计有电话的数量
    with_phone = sum(1 for c in companies if c.get("phone"))

    wb = Workbook()
    ws = wb.active
    ws.title = "潜在客户"

    # === 样式 ===
    header_font = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    data_font = Font(name="微软雅黑", size=10)
    data_align = Alignment(vertical="center", wrap_text=True)
    phone_font = Font(name="微软雅黑", size=10, bold=True, color="C00000")

    thin_border = Border(
        left=Side(style="thin", color="B4C6E7"),
        right=Side(style="thin", color="B4C6E7"),
        top=Side(style="thin", color="B4C6E7"),
        bottom=Side(style="thin", color="B4C6E7"),
    )
    even_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    phone_highlight = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    # === 标题行 ===
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
    title_cell = ws.cell(row=1, column=1)
    title_cell.value = (
        f"布匹中介 - 潜在客户名单  |  "
        f"共 {len(companies)} 家  |  "
        f"有电话 {with_phone} 家  |  "
        f"采集时间 {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    )
    title_cell.font = Font(name="微软雅黑", bold=True, size=13, color="1F4E79")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    # === 表头 ===
    for col_idx, (header_name, width, _) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=2, column=col_idx, value=header_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[2].height = 24

    # === 数据行 ===
    # 排序：有电话的排前面
    companies.sort(key=lambda c: (0 if c.get("phone") else 1, c.get("city", "")))

    for row_idx, company in enumerate(companies, 1):
        excel_row = row_idx + 2
        has_phone = bool(company.get("phone"))

        for col_idx, (_, _, field) in enumerate(COLUMNS, 1):
            value = row_idx if field is None else company.get(field, "")
            cell = ws.cell(row=excel_row, column=col_idx, value=value)
            cell.border = thin_border
            cell.alignment = data_align

            # 电话列高亮
            if field == "phone" and has_phone:
                cell.font = phone_font
                cell.fill = phone_highlight
            else:
                cell.font = data_font

            # 斑马纹（无电话高亮时）
            if row_idx % 2 == 0 and not (field == "phone" and has_phone):
                cell.fill = even_fill

    # === 冻结和筛选 ===
    ws.freeze_panes = "B3"
    last_col = get_column_letter(len(COLUMNS))
    ws.auto_filter.ref = f"A2:{last_col}{len(companies) + 2}"

    # === 统计 sheet ===
    _write_stats(wb.create_sheet("统计概览"), companies)

    wb.save(filepath)
    logger.info(f"已导出到: {filepath}")
    return filepath


def _write_stats(ws, companies):
    """统计概览页"""
    bold = Font(name="微软雅黑", bold=True, size=11)
    normal = Font(name="微软雅黑", size=10)
    title_font = Font(name="微软雅黑", bold=True, size=14, color="1F4E79")
    header_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 12

    with_phone = sum(1 for c in companies if c.get("phone"))
    rows = [
        (title_font, "采集统计概览", "", ""),
        (normal, "", "", ""),
        (bold, "指标", "数量", ""),
        (normal, "企业总数", len(companies), ""),
        (normal, "有联系电话", with_phone, ""),
        (normal, "无联系电话", len(companies) - with_phone, ""),
        (normal, "电话覆盖率", f"{with_phone/max(len(companies),1)*100:.1f}%", ""),
        (normal, "", "", ""),
        (bold, "城市", "企业数", "有电话"),
    ]

    # 按城市统计
    city_stats = {}
    for c in companies:
        city = c.get("city", "未知")
        if city not in city_stats:
            city_stats[city] = {"total": 0, "with_phone": 0}
        city_stats[city]["total"] += 1
        if c.get("phone"):
            city_stats[city]["with_phone"] += 1

    for city, stats in sorted(city_stats.items(), key=lambda x: -x[1]["total"]):
        rows.append((normal, city, stats["total"], stats["with_phone"]))

    for row_idx, (font, *values) in enumerate(rows, 1):
        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font
            if font == bold:
                cell.fill = header_fill
